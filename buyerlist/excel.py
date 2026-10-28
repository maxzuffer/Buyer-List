"""Stage 8 - the Excel in the reference longlist format, and its read-back validation.

Main sheets follow the reference longlist exactly: same column names and order, navy
header (00274D, white bold Calibri 9), red Feedback header (C00000), Excel
table TableStyleMedium2, duplicate-value highlight on Buyer, no gridlines,
same column widths, rows sorted Ranking -> Country -> Buyer. Feedback carries
only the analyst's recorded feedback from earlier review rounds; Status stays
empty for the deal team. Backup sheets carry the scoring, the PE reroutes, the
exclusions, the rounds and feedback, and the method.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import rounds as R
from .assemble import PLATFORM_MARK, load_list
from .common import OUTPUT, Issues, is_url, now_iso, read_json, rel, sha256_file, text, write_json
from .figures import bounds, display
from .mandate import load_mandate
from .scoring import platform_rule_text, platform_rules

STRATEGIC_HEADERS = ["Ranking", "Buyer", "Country", "Bereich", "Description", "Website", "Fiscal Year",
                     "Revenues (€m)", "EBITDA", "Employees", "Shareholder", "Comments", "Feedback", "Status"]
FINANCIAL_HEADERS = ["Ranking", "Buyer", "Country", "Portfolio Company", "Description", "Website",
                     "Feedback", "Status"]
STRATEGIC_WIDTHS = [13.36, 31.63, 9.45, 51.54, 47.45, 47.45, 12, 33.82, 36.54, 38.63, 60.45, 38.63, 19, 8.63]
FINANCIAL_WIDTHS = [13.36, 31.63, 11.18, 61, 61, 64.09, 44.82, 18.63]
NAVY, RED, BLUE = "FF00274D", "FFC00000", "FFDDEBF7"
TABLE_STYLE = "TableStyleMedium2"
SHEETS = ["Strategic", "Financial", "Strategic scoring", "Financial scoring", "PE reroutes", "Excluded",
          "Reserve", "Rounds & feedback", "Method"]

HEAD_FONT = Font(name="Calibri", size=9, bold=True, color="FFFFFFFF")
BODY_FONT = Font(name="Calibri", size=11, color="FF000000")
LINK_FONT = Font(name="Calibri", size=11, underline="single", color="FF0563C1")
TOP = Alignment(horizontal="left", vertical="top")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
HEAD_AL = Alignment(vertical="center", wrap_text=True)


def xlsx_path(run: Path) -> Path:
    """The final list keeps the plain name; a review round adds ' - round n of N'."""
    from .rounds import info
    m = load_mandate(run)
    r = info(run, m)
    tail = "" if r["final"] else f" - round {r['no']} of {r['of']}"
    return OUTPUT / f"{m['as_of'].replace('-', '')}_Longlist {m['project']}{tail}.xlsx"


def lock_path(run: Path) -> Path:
    return run / "08_final.lock"


def _header(ws, headers, widths, red_col=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = HEAD_FONT
        c.alignment = HEAD_AL
        c.fill = PatternFill("solid", fgColor=RED if h == red_col else NAVY)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def _table(ws, name, ncols, nrows):
    last = ws.cell(row=1, column=ncols).column_letter
    t = Table(displayName=name, ref=f"A1:{last}{max(nrows + 1, 2)}")
    t.tableStyleInfo = TableStyleInfo(name=TABLE_STYLE, showFirstColumn=False, showLastColumn=False,
                                      showRowStripes=True, showColumnStripes=False)
    ws.add_table(t)


def _dup_rule(ws, nrows):
    dxf = DifferentialStyle(font=Font(color="FF9C0006"), fill=PatternFill(bgColor="FFFFC7CE"))
    ws.conditional_formatting.add(f"B2:B{max(nrows + 1, 2)}", Rule(type="duplicateValues", dxf=dxf))


def _view(ws, zoom):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom


def _cell(ws, row, col, value, wrap=False, numfmt=None, link=None, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = LINK_FONT if link else BODY_FONT
    c.alignment = WRAP if wrap else TOP
    if numfmt and isinstance(value, (int, float)):
        c.number_format = numfmt
    if link:
        c.hyperlink = link
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def rounds_text(d: dict) -> str:
    r = d.get("round") or {}
    s = r.get("schedule") or []
    if not s:
        return "single round (no analyst review rounds)"
    return (f"round {r['no']} of {r['of']}; new names per side and round: {', '.join(str(x) for x in s)} (the last "
            "round fills the list to the target); the analyst's feedback on every new name is required before "
            "the next round")


def sort_key(r):
    return (r["rank"], text(r.get("country")), text(r.get("name")).casefold())


def _backup_sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    _header(ws, headers, widths or [18] * len(headers))
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            _cell(ws, i, j, v, wrap=isinstance(v, str) and len(v) > 40)
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(headers)).column_letter}{max(len(rows) + 1, 2)}"
    ws.freeze_panes = "C2"
    _view(ws, 80)
    return ws


def _urls(srcs):
    return "\n".join(s.get("url", "") for s in srcs or [])


def money(fig):
    """(cell value, number format) for a revenue / EBITDA figure."""
    if isinstance(fig, dict) and str(fig.get("status", "")).upper() == "VERIFIED" and isinstance(fig.get("value"), (int, float)):
        v = round(float(fig["value"]), 1)
        return v, ("#,##0" if abs(v) >= 100 else "#,##0.0")
    return display(fig), None


def strategic_cells(r: dict) -> list:
    """[(value, kind, numfmt)] for one Strategic row; kind: text|wrap|link|num."""
    rev, rf = money(r.get("revenue"))
    eb, ef = money(r.get("ebitda"))
    emp = display(r.get("employees"))
    return [(r["rank"], "num", None), (r["name"], "text", None), (r["country"], "text", None),
            (r["bereich"], "text", None), (r["description"], "wrap", None), (r["website"], "link", None),
            (r["fiscal_year"], "text", None), (rev, "num", rf), (eb, "num", ef),
            (emp, "num", "#,##0" if isinstance(emp, (int, float)) else None),
            (r["shareholders"], "wrap", None), (r["comments"], "wrap", None),
            (r.get("feedback_text") or None, "wrap", None), (None, "text", None)]


def financial_cells(r: dict) -> list:
    return [(r["rank"], "num", None), (r["name"], "text", None), (r["country"], "text", None),
            (r["portfolio_company"] or None, "wrap", None), (r["description"], "wrap", None),
            (r["website"], "link", None), (r.get("feedback_text") or None, "wrap", None), (None, "text", None)]


def _put(ws, row, cells, fill=None):
    for col, (v, kind, nf) in enumerate(cells, 1):
        _cell(ws, row, col, v, wrap=kind == "wrap", numfmt=nf, link=v if kind == "link" else None, fill=fill)


def write(run: Path) -> Path:
    d = load_list(run)
    if "round" not in d:
        from .common import StageError
        raise StageError("06_list_data.json was written before the PE-platform / rounds change - rebuild it first: "
                         "python -m buyerlist reopen 5e --reason \"...\", debate collect, gate 5e, assemble, gate 6, "
                         "audit, then excel")
    m = d["mandate"]
    wb = Workbook()
    # --- Strategic
    ws = wb.active
    ws.title = "Strategic"
    _header(ws, STRATEGIC_HEADERS, STRATEGIC_WIDTHS, red_col="Feedback")
    rows = sorted(d["strategic"]["selected"], key=sort_key)
    for i, r in enumerate(rows, 2):
        _put(ws, i, strategic_cells(r))
    _table(ws, "Strategic", len(STRATEGIC_HEADERS), len(rows))
    _dup_rule(ws, len(rows))
    _view(ws, 60)
    # --- Financial
    wf = wb.create_sheet("Financial")
    _header(wf, FINANCIAL_HEADERS, FINANCIAL_WIDTHS, red_col="Feedback")
    frows = sorted(d["financial"]["selected"], key=sort_key)
    for i, r in enumerate(frows, 2):
        _put(wf, i, financial_cells(r), fill=BLUE if r["platforms"] else None)
    _table(wf, "Financial", len(FINANCIAL_HEADERS), len(frows))
    _dup_rule(wf, len(frows))
    _view(wf, 67)
    # --- backups
    sw = m["weights"]["strategic"]
    _backup_sheet(wb, "Strategic scoring",
                  ["Ranking", "Buyer", "ID", "List", "Fit", "Confidence", "Product", "Customer", "Geography",
                   "Value chain", "Synergy", "Role", "For (judge)", "Against (judge)", "Revenue gate",
                   "Revenue basis", "Ownership type", "M&A activity", "Financial capability", "Sources"],
                  [[r["rank"], r["name"], r["id"], lst, r["fit"], r["confidence"], r["scores"]["product"],
                    r["scores"]["customer"], r["scores"]["geography"], r["scores"]["value_chain"],
                    r["scores"]["synergy"], r["role"], r["for"], r["against"], r["revenue_gate"], r["revenue_note"],
                    r["ownership_type"], r["ma_text"], f"{r['capability']['rating']}: {r['capability']['justification']}",
                    _urls(r["deciding_sources"] + r["against_sources"])]
                   for lst, src in (("Selected", d["strategic"]["selected"]), ("Reserve", d["strategic"]["reserve"]))
                   for r in src],
                  [9, 30, 8, 9, 6, 11, 8, 8, 9, 9, 8, 11, 50, 50, 11, 40, 14, 40, 40, 50])
    _backup_sheet(wb, "Financial scoring",
                  ["Ranking", "Buyer", "ID", "List", "Rank basis", "Investor fit", "Own fit (T/G/C)", "Investor rank",
                   "Platform rank", "Confidence", "Ticket", "Geography", "Criteria", "Portfolio", "Ticket (investor)",
                   "Geographic focus", "Sector focus", "Deployment", "For (judge)", "Against (judge)",
                   "PE-owned platforms", "Sources"],
                  [[r["rank"], r["name"], r["id"], lst, r["rank_basis"], r["fit"], r["own_fit"], r["fi_rank"],
                    r["platform_rank"],
                    r["confidence"], r["scores"]["ticket"], r["scores"]["geography"], r["scores"]["criteria"],
                    r["scores"]["portfolio"], r["ticket"], r["geo_focus"], r["sector_focus"], r["deployment"],
                    r["for"], r["against"], "; ".join(p["name"] for p in r["platforms"]),
                    _urls(r["deciding_sources"] + r["against_sources"])]
                   for lst, src in (("Selected", d["financial"]["selected"]), ("Reserve", d["financial"]["reserve"]))
                   for r in src],
                  [9, 30, 8, 9, 11, 8, 9, 8, 8, 11, 7, 9, 8, 9, 35, 35, 35, 11, 50, 50, 35, 50])
    sel_f = {r["id"] for r in d["financial"]["selected"]}
    _backup_sheet(wb, "PE reroutes",
                  ["Strategic (PE-owned)", "Bereich", "Strategic fit", "Strategic rank", "Product", "Section",
                   "Confidence", "Sponsor(s)", "Platform status", "Qualifies for", "Reason", "Shown on Financial list with",
                   "Sponsor(s) listed without the platform"],
                  [[p["name"], p["bereich"], p["fit"], p["rank"], p["product"], p["section"], p["confidence"],
                    "; ".join(p["sponsor_names"]), p["status"],
                    "; ".join(dict(zip(p["sponsors"], p["sponsor_names"]))[s] for s in p["listed_with"]),
                    p["reason"],
                    "; ".join(dict(zip(p["sponsors"], p["sponsor_names"]))[s] for s in p.get("shown_with", [])),
                    "; ".join(dict(zip(p["sponsors"], p["sponsor_names"]))[s] for s in dict.fromkeys(p["sponsors"])
                              if s in sel_f and s not in p["listed_with"])]
                   for p in d["rerouted"]], [35, 30, 10, 10, 9, 12, 11, 40, 16, 35, 70, 30, 30])
    _backup_sheet(wb, "Excluded", ["Side", "Buyer", "ID", "Section", "Reason", "Website"],
                  [[side, r["name"], r["id"], r["section"], r.get("reason", ""), r.get("website", "")]
                   for side in ("strategic", "financial") for r in d[side]["excluded"]],
                  [10, 35, 8, 16, 70, 40])
    _backup_sheet(wb, "Reserve", ["Side", "Ranking", "Buyer", "ID", "Fit", "Confidence", "Website"],
                  [[side, r["rank"], r["name"], r["id"], r["fit"], r["confidence"], r["website"]]
                   for side in ("strategic", "financial") for r in d[side]["reserve"]],
                  [10, 9, 35, 8, 6, 11, 40])
    audit = run / "07_audit" / "summary.json"
    au = read_json(audit) if audit.exists() else {}
    ev, eq = m["ev"]["value"], m["equity_ticket"]
    fw = m["weights"]["financial"]
    method = [
        ["Project", m["project"]], ["Client", m["client"]["name"]], ["As of", m["as_of"]],
        ["Revenue (assumption)", f"EUR {m['revenue']['value']:g}m ({m['revenue']['method']})"],
        ["EBITDA (assumption)", f"EUR {m['ebitda']['value']:g}m"],
        ["Enterprise value", f"EUR {ev[0]:g}-{ev[1]:g}m ({m['ev']['method']})"],
        ["Equity cheque", f"EUR {eq[0]:g}-{eq[1]:g}m"], ["Transaction", m["deal_type"]],
        ["Strategic revenue gate", f"strict: proven revenue > EUR {m['revenue_gate']['min_revenue']:g}m "
                                   "(VERIFIED, or ESTIMATE whose low end is above it); unknown = excluded"],
        ["PE reroute", "every PE-owned strategic leaves the Strategic sheet (sponsors together > 50% or one "
                       "sponsor in control)"],
        ["PE platform rule", platform_rule_text(platform_rules(m)) + f"; light-blue rows on the Financial sheet, "
                             f"'{PLATFORM_MARK}' in Portfolio Company"],
        ["Investor holdings", "only holdings with the same products or the same product category as the client "
                              "count (Portfolio score, 'Relevant holdings'); adjacent businesses do not"],
        ["Strategic fit weights", ", ".join(f"{k} {v:.0%}" for k, v in sw.items())],
        ["Financial fit weights", ", ".join(f"{k} {v:.0%}" for k, v in fw.items())],
        ["Ranks", "1: fit >= 3.8 | 2: 3.0-3.7 | 3: 2.0-2.9 | below 2.0: not listed. Financial rank = better of "
                  "investor rank and the rank of a qualifying platform. Strategic: vertical buyers and product+region "
                  "<= 2 not listed."],
        ["Debate", "advocate + critic research independently (verbatim quotes), program + Controller check every "
                   "quote, one retry round, judge sees anonymised Brief A/B (Business-Dev-2.0 bd-fit-debate)"],
        ["Debate points strategic", str(d["debate"]["strategic"].get("points"))],
        ["Debate points financial", str(d["debate"]["financial"].get("points"))],
        ["Audit", str(au) if au else "n/a"],
        ["Sort order (main sheets)", "Ranking, Country, Buyer (as in the reference longlist)"],
        ["Rounds", rounds_text(d)],
        ["Feedback / Status", "Feedback: the analyst's recorded feedback from earlier rounds (R<n> <status>: "
                              "<analyst's words>); otherwise empty for the deal team. Status: empty for the deal team"],
    ]
    fb_rows = []
    for rd in R.closed(run):
        f = R.round_feedback(run, rd["round"])
        fb_rows += [[rd["round"], x["side"], x["name"], x["id"], x["status"], x["note"], x["at"], x["by"]]
                    for x in f["items"]]
        fb_rows += [[rd["round"], x["type"], x["name"], "(added)", x["status"], x["note"], x["at"], x["by"]]
                    for x in f["new"]]
    _backup_sheet(wb, "Rounds & feedback",
                  ["Round", "Side", "Buyer", "ID", "Status", "Analyst feedback", "Recorded", "By"],
                  fb_rows or [["-", "", "no analyst feedback recorded yet", "", "", "", "", ""]],
                  [8, 11, 35, 9, 11, 80, 22, 16])
    _backup_sheet(wb, "Method", ["Item", "Value"], method, [28, 120])
    out = xlsx_path(run)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        out.unlink()
    wb.save(out)
    return out


# ---------------------------------------------------------------------------
# Validation (gate 8)
# ---------------------------------------------------------------------------
def _fill(c) -> str:
    try:
        return (c.fill.fgColor.rgb or "").upper()
    except Exception:
        return ""


def _same(a, b) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) < 1e-6
    return (a if a not in ("",) else None) == (b if b not in ("",) else None)


def gate8(run: Path) -> tuple:
    iss = Issues()
    out = xlsx_path(run)
    if not out.exists():
        iss.error("excel", rel(out), "missing - run: python -m buyerlist excel")
        return iss, []
    d = load_list(run)
    m = d["mandate"]
    from .profiles import revenue_gate
    thr = m["revenue_gate"]["min_revenue"]
    wb = load_workbook(out)
    if wb.sheetnames[:2] != ["Strategic", "Financial"]:
        iss.error("excel", "sheets", f"first sheets must be Strategic, Financial - got {wb.sheetnames[:2]}")
        return iss, [out]
    s_keys = set()
    for name, headers, side, builder in (("Strategic", STRATEGIC_HEADERS, "strategic", strategic_cells),
                                         ("Financial", FINANCIAL_HEADERS, "financial", financial_cells)):
        ws = wb[name]
        got = [ws.cell(row=1, column=i).value for i in range(1, len(headers) + 1)]
        if got != headers:
            iss.error("excel", name, f"headers {got} != reference {headers}")
        for i, h in enumerate(headers, 1):
            want = RED if h == "Feedback" else NAVY
            if _fill(ws.cell(row=1, column=i)) != want:
                iss.error("excel", name, f"header {h} fill {_fill(ws.cell(row=1, column=i))} != {want}")
        if name not in ws.tables or ws.tables[name].tableStyleInfo.name != TABLE_STYLE:
            iss.error("excel", name, f"Excel table with {TABLE_STYLE} missing")
        if not any(r.type == "duplicateValues" for cf in ws.conditional_formatting for r in cf.rules):
            iss.error("excel", name, "duplicate-value highlight on Buyer missing")
        if ws.sheet_view.showGridLines:
            iss.error("excel", name, "gridlines must be off")
        rows = sorted(d[side]["selected"], key=sort_key)
        if (d.get("round") or {"final": True})["final"] and ws.max_row - 1 != m["targets"][side]:
            iss.error("excel", name, f"{ws.max_row - 1} rows, {m['targets'][side]} required")
        if ws.max_row - 1 != len(rows):
            iss.error("excel", name, f"{ws.max_row - 1} rows but list_data has {len(rows)} selected")
        keys = set()
        for i, r in enumerate(rows, 2):
            who = f"{name} row {i} {r['name']}"
            for col, (v, kind, nf) in enumerate(builder(r), 1):
                c = ws.cell(row=i, column=col)
                if not _same(c.value, v):
                    iss.error("excel", who, f"{headers[col - 1]} is {c.value!r}, list data says {v!r} - run excel again")
                if kind == "link" and (c.hyperlink is None or c.hyperlink.target != v or not is_url(v)):
                    iss.error("excel", who, "website missing, not a valid link, or different from the list data")
            from .universe import registrable_domain
            k = registrable_domain(r.get("website") or "") or r["id"]
            if k in keys or (side == "financial" and k in s_keys):
                iss.error("excel", who, "duplicate buyer (within or across sheets)")
            keys.add(k)
            if side == "strategic":
                if revenue_gate(r, thr)[0] != "PASS":
                    iss.error("excel", who, f"revenue does not prove > EUR {thr:g}m")
                if r["pe_owned"]:
                    iss.error("excel", who, "PE-owned strategic on the Strategic sheet")
                if not text(r["comments"]).startswith("Fit "):
                    iss.error("excel", who, "Comments must start with the adjudicated fit")
            else:
                pc = text(r["portfolio_company"])
                if r["platforms"]:
                    if not pc.startswith(PLATFORM_MARK):
                        iss.error("excel", who, f"PE platform label '{PLATFORM_MARK}' missing in Portfolio Company")
                    if _fill(ws.cell(row=i, column=1)) != BLUE:
                        iss.error("excel", who, "rerouted row must be light blue")
                elif PLATFORM_MARK in pc or _fill(ws.cell(row=i, column=1)) == BLUE:
                    iss.error("excel", who, "platform label / blue fill on a row without a rerouted platform")
        if side == "strategic":
            s_keys = keys
    for sname in SHEETS[2:]:
        if sname == "Rounds & feedback" and "round" not in d:        # list data written before rounds existed
            continue
        if sname not in wb.sheetnames:
            iss.error("excel", sname, "backup sheet missing")
    return iss, [out]


def write_lock(run: Path, summary: str) -> dict:
    out = xlsx_path(run)
    lock = {"xlsx": rel(out), "sha256": sha256_file(out), "validated_at": now_iso(), "summary": summary}
    write_json(lock_path(run), lock)
    return lock


def lock_ok(run: Path) -> tuple:
    lp = lock_path(run)
    if not lp.exists():
        return False, "no final.lock - run: python -m buyerlist validate"
    lock = read_json(lp)
    from .common import ROOT
    from . import state
    st = state.load(run)
    if state.next_stage(st) is not None:
        return False, f"stage {state.next_stage(st)} is not passed"
    if lock.get("xlsx") != rel(xlsx_path(run)):
        return False, f"lock is for {lock.get('xlsx')}, not {rel(xlsx_path(run))}"
    p = ROOT / lock["xlsx"]
    if not p.exists():
        return False, f"{lock['xlsx']} missing"
    digest = sha256_file(p)
    if digest != lock["sha256"] or (st["stages"]["8"].get("outputs") or {}).get(lock["xlsx"]) != digest:
        return False, f"{lock['xlsx']} changed after validation"
    return True, lock["summary"]
