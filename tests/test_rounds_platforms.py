"""PE platform rule, close holdings, review rounds and analyst feedback (user decisions 2026-10-28)."""
import json
import tempfile
import time
import unittest
from datetime import date
from pathlib import Path

from tests import fixture as F
from tests.test_hooks import hook
from buyerlist import assemble, audit, excel, fiscore, normalize, platforms, rounds, route, scoring, state, universe
from buyerlist.common import StageError, read_json, write_json
from buyerlist.profiles import validate_financial


class PlatformRule(unittest.TestCase):
    R = scoring.platform_rules()

    def check(self, fit, product, sponsor_fit, section="ranked"):
        return scoring.platform_check(fit, product, section, sponsor_fit, self.R)[0]

    def test_own_fit_leaves_portfolio_out(self):
        w = scoring.FINANCIAL_WEIGHTS
        self.assertEqual(scoring.own_fit({"ticket": 4, "geography": 4, "criteria": 4, "portfolio": 1}, w), 4.0)
        self.assertEqual(scoring.own_fit({"ticket": 2, "geography": 4, "criteria": 3, "portfolio": 5}, w), 2.9)
        self.assertIsNone(scoring.own_fit({"ticket": 2}, w))

    def test_general_threshold(self):
        self.assertTrue(self.check(3.0, 4, 3.5))
        self.assertFalse(self.check(3.0, 3, 3.5))               # product 3: no close product overlap
        self.assertFalse(self.check(2.9, 4, 3.5))               # below the Rank-2 floor
        self.assertFalse(self.check(4.5, 5, 4.0, "vertical"))
        self.assertFalse(self.check(4.5, 5, 4.0, "leermeldung"))

    def test_weak_sponsor_needs_rank1_same_products(self):
        self.assertFalse(self.check(3.8, 4, 2.9))
        self.assertFalse(self.check(3.7, 5, 2.9))
        self.assertTrue(self.check(3.8, 5, 2.9))
        self.assertTrue(self.check(3.8, 5, None))               # no own investor score counts as weak
        self.assertTrue(self.check(3.0, 4, 3.0))                # 3.0 is not below the floor
        ok, why = scoring.platform_check(3.8, 4, "ranked", 2.8, self.R)
        self.assertIn("own fit 2.8 < 3.0", why)
        self.assertIn("product 4 < 5", why)

    def test_mandate_override(self):
        r = scoring.platform_rules({"platform_rule": {"min_fit": 3.5}})
        self.assertEqual(r["min_fit"], 3.5)
        self.assertEqual(r["strict_fit"], 3.8)


class Holdings(unittest.TestCase):
    def test_overlap_normalised(self):
        self.assertEqual(normalize.overlap("Same product category"), "same category")
        self.assertEqual(normalize.overlap("same_products"), "same products")
        self.assertEqual(normalize.overlap(5), "same products")
        self.assertEqual(normalize.overlap("Complementary (software)"), "adjacent")
        p = normalize.profile({"relevant_portfolio": [{"name": "A", "overlap": "Same Category"}]})
        self.assertEqual(p["relevant_portfolio"][0]["overlap"], "same category")

    def test_overlap_validated(self):
        cand = {"id": "f-009", "name": "X Partners", "type": "financial"}
        p = F.f_profile("f-009", "X Partners")
        p["relevant_portfolio"] = [{"name": "NoClass", "note": "dc"},
                                   {"name": "NoSource", "overlap": "same category", "note": "dc"},
                                   {"name": "Soft", "overlap": "adjacent", "note": "software"}]
        errs = [i.message for i in validate_financial(p, cand, date(2026, 10, 28)).errors]
        self.assertTrue(any("NoClass: overlap must be" in e for e in errs), errs)
        self.assertTrue(any("NoSource: a same-products" in e for e in errs), errs)
        self.assertFalse(any(e.startswith("Soft") for e in errs), errs)

    def test_close_holdings_and_portfolio(self):
        p = {"relevant_portfolio": [{"name": "IRIDEOS Retelit", "overlap": "same category"},
                                    {"name": "SoftCo", "overlap": "adjacent"},
                                    {"name": "FiberNet", "overlap": "same products", "website": "https://fibernet.example"}]}
        plats = [{"name": "Retelit S.p.A.", "website": "https://retelit.it"}]
        self.assertEqual([h["name"] for h in platforms.close_holdings(p, plats)], ["FiberNet"])
        self.assertFalse(platforms.same_company({"name": "Data Group"}, {"name": "Group Holding"}))
        self.assertEqual(fiscore.portfolio(p, ["Retelit"], plats)[0], 5)
        self.assertEqual(fiscore.portfolio(p, [], plats)[0], 3)            # the platform does not count as a holding
        self.assertEqual(fiscore.portfolio({"relevant_portfolio": [{"name": "SoftCo", "overlap": "adjacent"}]},
                                           [], [])[0], 2)


def transcript(text, ts, cmd="validate"):
    f = tempfile.NamedTemporaryFile("w", suffix=".jsonl", delete=False)
    f.write(json.dumps({"type": "user", "timestamp": ts, "message": {"content": "next"}}) + "\n")
    f.write(json.dumps({"type": "assistant", "timestamp": ts, "message": {"content": [
        {"type": "tool_use", "name": "Bash", "id": "t1", "input": {"command": f".venv/bin/python -m buyerlist {cmd}"}}]}}) + "\n")
    f.write(json.dumps({"type": "assistant", "timestamp": ts, "message": {"content": [{"type": "text", "text": text}]}}) + "\n")
    f.close()
    return f.name


def weak(p):
    """Turn an investor profile into one whose own fit is below 3.0 (ticket n/a, US-only, real estate)."""
    p = dict(p, hq_country="US", offices=["New York"])
    p["ticket"] = {"status": "n/a", "reason": "not stated"}
    p["latest_fund"] = None
    p["geo_focus"] = {"summary": "North America", "quote": "North America", "source": "https://x.example"}
    p["sector_focus"] = {"summary": "real estate", "quote": "real estate", "source": "https://x.example"}
    return p


class Rounds(unittest.TestCase):
    def setUp(self):
        self.guard = F.ActiveGuard().__enter__()
        self.saved = rounds.MIN_EXTRA
        rounds.MIN_EXTRA = {"strategic": 0, "financial": 0}       # the fixture universe is tiny

    def tearDown(self):
        rounds.MIN_EXTRA = self.saved
        self.guard.__exit__()

    def test_schedule(self):
        m = {"targets": {"strategic": 100, "financial": 100}, "rounds": [10, 20, 20, 50]}
        self.assertEqual(rounds.validate_schedule(m), [])
        self.assertTrue(rounds.validate_schedule(dict(m, rounds=[60, 40, 10])))    # 100 before the last round
        self.assertTrue(rounds.validate_schedule(dict(m, rounds=[0, 5])))
        self.assertEqual(rounds.cumulative(m, 2), {"strategic": 30, "financial": 30})
        self.assertEqual(rounds.cumulative(m, 4), {"strategic": 100, "financial": 100})
        self.assertEqual(rounds.validate_schedule({"targets": m["targets"]}), [])     # no schedule: one round
        self.assertEqual(rounds.default_schedule({"strategic": 100, "financial": 100}), [10, 20, 20, 50])
        self.assertEqual(rounds.default_schedule({"strategic": 30, "financial": 30}), [3, 6, 6, 15])
        self.assertEqual(rounds.default_schedule({"strategic": 3, "financial": 3}), [1, 1, 1])
        from buyerlist.cli import _parse_rounds
        self.assertIsNone(_parse_rounds("off", [100, 100]))
        with self.assertRaises(StageError):
            _parse_rounds("10;20", [100, 100])

    def test_two_rounds_with_feedback(self):
        run = F.build("8", rounds=[2, 1])
        d = assemble.load_list(run)
        self.assertEqual((d["round"]["no"], d["round"]["final"]), (1, False))
        s_ids = sorted(r["id"] for r in d["strategic"]["selected"])
        f_ids = sorted(r["id"] for r in d["financial"]["selected"])
        self.assertEqual(s_ids, ["s-001", "s-002"])                               # round quota 2, best first
        self.assertEqual(f_ids, ["f-004", "f-r001"])
        self.assertTrue(excel.xlsx_path(run).name.endswith(" - round 1 of 2.xlsx"))
        self.assertTrue(excel.lock_ok(run)[0])
        from openpyxl import load_workbook
        self.assertIn("Rounds & feedback", load_workbook(excel.xlsx_path(run)).sheetnames)
        sampled = {(x["kind"], x["id"]) for x in audit.sample(run)}
        for i in s_ids:
            self.assertIn(("strategic", i), sampled)                               # every new row audited
        for i in f_ids:
            self.assertIn(("financial", i), sampled)

        # the Stop hook makes the finished round ask for feedback
        ts = time.strftime("%Y-%m-%dT%H:%M:%S", time.gmtime(time.time() - 600)) + "Z"
        res = hook("require_pipeline_complete.py", {"transcript_path": transcript("Validation: PASS - 0 errors", ts)})
        self.assertEqual(res.get("decision"), "block")
        self.assertIn("Feedback needed:", res["reason"])
        res = hook("require_pipeline_complete.py", {"transcript_path": transcript(
            "Validation: PASS - 0 errors\nFeedback needed: round 1 - BigTel AG, DataCo GmbH ...", ts)})
        self.assertEqual(res, {})

        # feedback: the analyst's words, every new name, notes for drop / data-wrong
        with self.assertRaises(StageError):
            rounds.add_feedback(run, "s-002", "drop", "")
        with self.assertRaises(StageError):
            rounds.add_feedback(run, "s-003", "keep", "x")                          # not on the list
        rounds.add_feedback(run, "s-001", "keep", "exactly the right peer")
        rounds.add_feedback(run, "s-002", "drop", "too small to care")
        rounds.add_feedback(run, "f-r001", "maybe", "")
        # feedback of the open round does not change the round's list (gate 6 still passes)
        self.assertTrue(assemble.gate6(run)[0].ok())
        with self.assertRaises(StageError):
            rounds.next_round(run, "x")                                            # f-004 has no feedback
        self.assertEqual([x[1] for x in rounds.missing(run)], ["f-004"])
        rounds.add_feedback(run, "f-004", "data-wrong", "Zeta's ticket is EUR 50-150m")
        rounds.add_name(run, "strategic", "NewCo AG", "https://newco.example", "DE", "keep", "obvious buyer, missing")
        # feedback complete: the Stop hook wants the next round (or PAUSED)
        fcmd = "feedback add --id f-004 --status data-wrong --note x"
        res = hook("require_pipeline_complete.py", {"transcript_path": transcript("Recorded.", ts, fcmd)})
        self.assertIn("round next", res.get("reason", ""))
        res = hook("require_pipeline_complete.py",
                   {"transcript_path": transcript("PAUSED: analyst asked to wait", ts, fcmd)})
        self.assertEqual(res, {})
        res = rounds.next_round(run, "analyst feedback round 1 recorded")
        self.assertEqual((res["round"], res["reprofile"]), (2, ["f-004"]))
        st = state.load(run)
        self.assertEqual(state.status(st, "2"), "open")
        self.assertTrue((universe.raw_dir(run) / f"AN-r{st['topup_round']}.json").exists())

        # the list of round 2 from the round-1 data: drop applied, keep and earlier names stay, final fill
        d2 = assemble.compute(run)
        self.assertTrue(d2["round"]["final"])
        s2 = {r["id"]: r for r in d2["strategic"]["selected"]}
        self.assertNotIn("s-002", s2)
        excl = {r["id"]: r for r in d2["strategic"]["excluded"]}
        self.assertEqual(excl["s-002"]["section"], "analyst_drop")
        self.assertIn("too small to care", excl["s-002"]["reason"])
        self.assertEqual(s2["s-001"]["feedback_text"], "R1 keep: exactly the right peer")
        self.assertEqual(s2["s-001"]["round_added"], 1)
        self.assertEqual(d2["shortfall"]["strategic"], 1)                         # s-002 dropped, NewCo not sourced yet
        self.assertEqual(excel.strategic_cells(s2["s-001"])[12][0], "R1 keep: exactly the right peer")
        f2 = {r["id"] for r in d2["financial"]["selected"]}
        self.assertTrue({"f-004", "f-r001"} <= f2)
        self.assertEqual(len(f2), 3)

        # round 2, stage 2: new lenses from the feedback, tasks that carry it, the analyst's name
        tr = st["topup_round"]
        with self.assertRaises(StageError):
            universe.sourcing_tasks(run, None, tr)                                  # no round-2 lenses yet
        L = read_json(universe.lenses_path(run))
        L["lenses"] += [{"id": f"{t[0].upper()}1{i}", "type": t, "title": "t", "min": 2, "round": 2,
                         "brief": "More owner-managed peers like the ones the analyst kept, as the analyst asked."}
                        for t in ("strategic", "financial") for i in (0, 1)]
        write_json(universe.lenses_path(run), L)
        idx = universe.sourcing_tasks(run, None, tr)
        self.assertEqual(sorted(t["id"] for t in idx), sorted(f"{x}-r{tr}" for x in ("S10", "S11", "F10", "F11")))
        body = Path(idx[0]["file"]).read_text()
        for s in ("Analyst feedback from the earlier rounds", "too small to care", "NewCo AG"):
            self.assertIn(s, body)
        fresh = {"name": "FreshDC GmbH", "type": "strategic", "website": "https://freshdc.example", "hq_country": "DE",
                 "hq_city": "Berlin", "segment": "DC", "rationale": "r", "source_url": "https://freshdc.example"}
        # Alpha Capital (so far only a route-added sponsor, f-r001) is now sourced: it gets a universe id
        alpha = {"name": "Alpha Capital", "type": "financial", "website": "https://alpha-capital.example",
                 "hq_country": "DE", "hq_city": "Munich", "segment": "PE", "rationale": "r",
                 "source_url": "https://alpha-capital.example"}
        for t in idx:
            write_json(Path(t["outputs"][0]), {"lens": t["id"].split("-")[0], "round": tr,
                                               "candidates": [fresh if t["id"].startswith("S") else alpha]})
        uni = universe.merge(run)
        new = {c["name"]: c["id"] for c in uni["candidates"] if c["name"] in ("NewCo AG", "FreshDC GmbH",
                                                                               "Alpha Capital")}
        self.assertEqual(len(new), 3)
        F.g(run, "2")
        # the analyst's maybe on f-r001 now belongs to the universe id (matched by website)
        self.assertIn(new["Alpha Capital"], rounds.shown(run)["financial"])
        self.assertNotIn("f-r001", rounds.shown(run)["financial"])
        # gate 2 round checks: too few round lenses, stale digest, missing analyst name
        lenses = read_json(universe.lenses_path(run))["lenses"]
        self.assertFalse(universe.round_checks(run, [l for l in lenses if l["id"] != "S11"]).ok())
        used = read_json(universe.feedback_used_path(run))
        write_json(universe.feedback_used_path(run), dict(used, digest="stale"))
        self.assertFalse(universe.round_checks(run, lenses).ok())
        write_json(universe.feedback_used_path(run), used)
        an = universe.raw_dir(run) / f"AN-r{tr}.json"
        keep = an.read_text()
        an.unlink()
        self.assertTrue(any("NewCo AG" in i.field for i in universe.round_checks(run, lenses).errors))
        an.write_text(keep)

        # stage 3: new names profiled; the analyst's data-wrong (f-004) must be re-profiled
        st = state.load(run); state.open_stage(run, st, "3"); state.save(run, st)
        rev = {"status": "VERIFIED", "value": 60, "fy": "2025", "scope": "group", "source": "AR"}
        for nm in ("NewCo AG", "FreshDC GmbH"):
            write_json(run / "03_profiles" / f"{new[nm]}.json", F.s_profile(new[nm], nm, rev, [], False))
        write_json(run / "03_profiles" / f"{new['Alpha Capital']}.json",
                   F.f_profile(new["Alpha Capital"], "Alpha Capital", "https://alpha-capital.example"))
        from buyerlist import gates
        iss, _ = gates.gate3(run)
        self.assertTrue(any("f-004" in i.message and "wrong data" in i.message for i in iss.errors), iss.render())
        write_json(run / "03_profiles" / "f-004.json", F.f_profile("f-004", "Zeta Infra"))
        F.g(run, "3")
        st = state.load(run); state.open_stage(run, st, "4"); state.save(run, st)
        r = route.build(run)
        self.assertEqual(r["counts"]["financial_added"], 0)                       # Alpha is now a sourced investor
        F.g(run, "4")

    def test_dropped_only_sponsor_and_intermediate_quota(self):
        run = F.build("8", rounds=[1, 1, 1])
        d = assemble.load_list(run)
        self.assertEqual([len(d[t]["selected"]) for t in ("strategic", "financial")], [1, 1])
        sid, fid = d["strategic"]["selected"][0]["id"], d["financial"]["selected"][0]["id"]
        self.assertEqual(fid, "f-004")                                            # Zeta: s-007's only qualifying sponsor
        rounds.add_feedback(run, sid, "maybe", "")
        rounds.add_feedback(run, fid, "drop", "we never work with Zeta Infra")
        rounds.next_round(run, "round 1 feedback")
        d2 = assemble.compute(run)
        self.assertFalse(d2["round"]["final"])
        sel_s = [r["id"] for r in d2["strategic"]["selected"]]
        self.assertEqual(len(sel_s), 2)                                           # carried + quota 1 new
        self.assertIn(sid, sel_s)
        self.assertEqual(len(d2["financial"]["selected"]), 1)                     # the dropped one is not carried
        pl = {p["strategic_id"]: p for p in d2["rerouted"]}["s-007"]
        self.assertEqual(pl["status"], "not shown - qualifying sponsor(s) dropped by the analyst")
        ex = {r["id"]: r for r in d2["strategic"]["excluded"]}
        self.assertIn("never work with Zeta Infra", ex["s-007"]["reason"])
        self.assertIn("s-007", {c["id"] for c in d2["feedback_conflicts"]})

    def test_profile_scoring_portfolio_follows_the_platform_rule(self):
        run = F.build("5")
        for fid, name in (("f-004", "Zeta Infra"), ("f-r001", "Alpha Capital")):
            p = F.f_profile(fid, name, "https://alpha-capital.example" if fid == "f-r001" else None)
            write_json(run / "03_profiles" / f"{fid}.json", weak(p))
        sc = fiscore.compute(run)["companies"]
        # Alpha (weak) keeps s-005 (Rank 1, product 5); Zeta (weak) loses s-007 (3.5, product 4) and s-008
        self.assertEqual(sc["f-r001"]["scores"]["portfolio"], 5)
        self.assertEqual(sc["f-004"]["scores"]["portfolio"], 2)

    def test_final_round_platform_overflow(self):
        run = F.build("6")
        from buyerlist.mandate import load_mandate
        m = load_mandate(run)
        m["targets"]["financial"] = 1
        orig = assemble.load_mandate
        assemble.load_mandate = lambda r: m
        try:
            data = assemble.compute(run)
        finally:
            assemble.load_mandate = orig
        self.assertEqual(data["must_include_overflow"], 1)                        # s-005 via Alpha, s-007 via Zeta


if __name__ == "__main__":
    unittest.main()
