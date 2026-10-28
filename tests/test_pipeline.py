import unittest

from tests import fixture as F
from buyerlist import assemble, excel, gates, state
from buyerlist.common import StageError, read_json, write_json
from openpyxl import load_workbook


class Pipeline(unittest.TestCase):
    def setUp(self):
        self.guard = F.ActiveGuard().__enter__()

    def tearDown(self):
        self.guard.__exit__()

    def test_order_enforced(self):
        run = F.build("0")
        ok, msg = state.can_start(run, "3")
        self.assertFalse(ok)
        self.assertIn("stage 1", msg)
        from buyerlist.cli import main
        self.assertEqual(main(["gate", "4"]), 2)   # BLOCKED: prerequisites

    def test_tamper_marks_stale(self):
        run = F.build("2")
        m = read_json(run / "00_mandate.json")
        m["revenue"]["value"] = 5
        write_json(run / "00_mandate.json", m)
        with self.assertRaises(StageError):
            st = state.load(run)
            state.require(run, st, "3")
        st = state.load(run)
        self.assertEqual(state.status(st, "0"), "stale")
        self.assertEqual(state.status(st, "2"), "stale")

    def test_full_run(self):
        run = F.build("8")
        d = assemble.load_list(run)
        s_ids = [r["id"] for r in d["strategic"]["selected"]]
        self.assertEqual(sorted(s_ids), ["s-001", "s-002", "s-003"])           # SmallCo revenue-excluded, Reseller vertical
        self.assertNotIn("s-005", s_ids)                                          # PE-owned rerouted
        f_sel = {r["id"]: r for r in d["financial"]["selected"]}
        alpha, zeta = f_sel["f-r001"], f_sel["f-004"]
        self.assertEqual(alpha["rank"], 1)                                        # weak sponsor ranked via its
        self.assertEqual(alpha["rank_basis"], "platform")                         # Rank-1, same-products platform
        self.assertTrue(alpha["portfolio_company"].startswith(assemble.PLATFORM_MARK))
        self.assertEqual([x["strategic_id"] for x in alpha["platforms"]], ["s-005"])  # not s-007 (product 4)
        self.assertEqual([x["strategic_id"] for x in zeta["platforms"]], ["s-007"])   # not s-008 (product 3)
        self.assertNotIn("Tau Platform", zeta["portfolio_company"])
        rer = {x["strategic_id"]: x for x in d["rerouted"]}
        self.assertEqual(rer["s-007"]["listed_with"], ["f-004"])
        self.assertIn("needs Rank 1 and the same products", rer["s-007"]["reason"])
        self.assertEqual(rer["s-008"]["status"], "not listed")
        self.assertIn("product 3 < 4", rer["s-008"]["reason"])
        excl = {r["id"]: r["section"] for r in d["strategic"]["excluded"]}
        self.assertEqual(excl.get("s-008"), "pe_platform_not_listed")
        self.assertNotIn("s-007", excl)
        self.assertEqual(excl.get("s-004"), "revenue_excluded")
        self.assertEqual(excl.get("s-006"), "vertical")
        sc = read_json(run / "05_debate" / "scores_strategic.json")["companies"]
        self.assertTrue(sc["s-003"]["retried"])
        self.assertEqual(sc["s-001"]["fit"], 4.4)                                 # 0.5*5+0.2*4+0.2*4+0.05*3+0.05*3
        ok, msg = excel.lock_ok(run)
        self.assertTrue(ok, msg)
        wb = load_workbook(excel.xlsx_path(run))
        ws = wb["Strategic"]
        self.assertEqual([c.value for c in ws[1]], excel.STRATEGIC_HEADERS)
        self.assertEqual(ws["A1"].fill.fgColor.rgb, "FF00274D")
        self.assertEqual(ws["M1"].fill.fgColor.rgb, "FFC00000")
        self.assertIn("Strategic", ws.tables)
        wf = wb["Financial"]
        self.assertEqual([c.value for c in wf[1]], excel.FINANCIAL_HEADERS)
        self.assertEqual(wf["A2"].fill.fgColor.rgb, excel.BLUE)
        self.assertEqual(state.load(run)["status"], "complete")
        # tamper with the Excel (needs chmod: passed outputs are read-only) -> lock no longer valid
        import os
        os.chmod(excel.xlsx_path(run), 0o644)
        wb.save(excel.xlsx_path(run))
        self.assertFalse(excel.lock_ok(run)[0])

    def test_shortfall_and_no_padding(self):
        run = F.build("6")
        from buyerlist.mandate import load_mandate
        m = load_mandate(run)
        m["targets"]["strategic"] = 5
        orig = assemble.load_mandate
        assemble.load_mandate = lambda r: m
        try:
            data = assemble.compute(run)
        finally:
            assemble.load_mandate = orig
        self.assertEqual(data["shortfall"]["strategic"], 2)
        # hand-padding the list (stage re-opened) is caught by the recompute
        st = state.load(run)
        state.open_stage(run, st, "6")
        state.save(run, st)
        d = assemble.load_list(run)
        d["strategic"]["selected"].append(dict(d["strategic"]["selected"][0], id="s-004", name="SmallCo GmbH"))
        write_json(assemble.list_path(run), d)
        iss, _ = gates.GATES["6"](run)
        self.assertFalse(iss.ok())
        self.assertTrue(any("does not match" in i.message for i in iss.errors))


if __name__ == "__main__":
    unittest.main()
